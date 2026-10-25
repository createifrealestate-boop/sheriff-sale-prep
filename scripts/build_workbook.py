"""Build (or rebuild) the per-auction prep workbook.

Usage:
  python3 -I build_workbook.py classified.json out.xlsx --sale-date 2026-11-02 \
      --source "November-Sale-List-Updated-10-20.pdf (printed 10/20/2026 3:12 PM)" \
      [--previous previous.xlsx]   # last version (e.g. the Google Sheet exported as .xlsx)

Keeps everything typed in the YOUR ENTRY columns, keyed by Sale ID, and lists what
changed since the previous version on the Overview tab.
"""
import argparse, json, datetime as dt
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: F401 (kept for future use)

REF = Path(__file__).resolve().parent.parent / "reference"

# ------------------------------------------------------------------ layout
# (header, key, width, kind)  kind: text | money | pct | wrap | manual-money | manual-text | formula-*
COLUMNS = [
    ("Sale ID", "sale_id", 10, "text"),
    ("Case #", "case_no", 14, "text"),
    ("Property address", "address", 26, "text"),
    ("City", "city", 14, "text"),
    ("ZIP", "zip", 7, "text"),
    ("Municipality", "municipality", 16, "text"),
    ("Parcel ID(s)", "parcels", 16, "text"),
    ("County parcel ID (portal format)", "portal_parcel", 22, "text"),
    ("Property type", "property_type", 16, "text"),
    ("Sale category", "sale_category", 20, "text"),
    ("Owner type", "owner_type", 14, "text"),
    ("Defendant / owner", "defendant", 22, "text"),
    ("Plaintiff", "plaintiff", 22, "text"),
    ("Plaintiff attorney", "plaintiff_attorney", 20, "text"),
    ("Cost & Tax bid", "cost_tax_bid", 12, "money"),
    ("Free & clear min bid", "fc_min_bid", 12, "money"),
    ("Opening bid to watch", "opening_bid", 12, "money"),
    ("Reassessed damages (judgment)", "reassessed_damages", 13, "money"),
    ("Boxes  Svs · 3129.2 · 3129.3 · OK", "boxes", 14, "text"),
    ("__STATUS__", "status_note", 34, "wrap"),
    ("Flags / title & bidding risks", "flags_text", 34, "wrap"),
    ("Sheriff comments (verbatim)", "comments", 44, "wrap"),
    ("Transfer tax municipality", "tax_municipality", 20, "text"),
    ("Transfer tax rate (total)", "=RATE", 10, "formula-pct"),
    ("YOUR ENTRY: Assessed value", "assessed_value", 13, "manual-money"),
    ("Est. transfer tax on sheriff deed", "=TAX", 13, "formula-money"),
    ("YOUR ENTRY: Max bid", "max_bid", 12, "manual-money"),
    ("YOUR ENTRY: My notes", "my_notes", 30, "manual-text"),
    ("YOUR ENTRY: Sold for", "sold_for", 12, "manual-money"),
    ("YOUR ENTRY: Bidder #", "bidder_no", 10, "manual-text"),
    ("First listed here", "first_seen", 11, "text"),
]
MANUAL_KEYS = [k for _, k, _, kind in COLUMNS if kind.startswith("manual")]
HEADER_ROW, FIRST = 4, 5
TT_FIRST = 14  # first municipality row on Transfer Tax

NAVY = PatternFill("solid", fgColor="1F3A5F")
YELLOW_H = PatternFill("solid", fgColor="B7791F")
YELLOW = PatternFill("solid", fgColor="FFF6D5")
GREEN = PatternFill("solid", fgColor="E6F4EA")
AMBER = PatternFill("solid", fgColor="FDF0DC")
GREY = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="D0D0D0")
WHITE_B = Font(bold=True, color="FFFFFF")
MONEY_FMT = '"$"#,##0.00'
PCT_FMT = "0.0%"


def col_of(key):
    for i, (_, k, _, _) in enumerate(COLUMNS, 1):
        if k == key:
            return get_column_letter(i)


def tick(b):
    return "✔" if b else "–"


# ------------------------------------------------------------------ read previous version
def read_previous(path):
    prev = {"manual": {}, "visible_ids": set(), "snapshot": {}, "first_seen": {}, "log": []}
    if not path or not Path(path).exists():
        return prev
    wb = openpyxl.load_workbook(path, data_only=True)
    # hidden store of every manual entry ever made (covers cases that dropped off)
    if "_Manual" in wb.sheetnames:
        ws = wb["_Manual"]
        hdr = [c.value for c in ws[1]]
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r and r[0]:
                prev["manual"][str(r[0])] = {h: v for h, v in zip(hdr[1:], r[1:]) if h}
    # visible tabs win: they hold the latest edits (including deliberately cleared cells)
    for name in ("OK to Sell", "Almost Ready"):
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        hdr = {}
        for c in ws[HEADER_ROW]:
            if c.value:
                hdr[str(c.value).strip()] = c.column - 1
        byh = {h: k for h, k, _, _ in COLUMNS}
        if "Sale ID" not in hdr:
            continue
        for r in ws.iter_rows(min_row=FIRST, values_only=True):
            sid = r[hdr["Sale ID"]] if hdr["Sale ID"] < len(r) else None
            if not sid:
                continue
            sid = str(sid)
            prev["visible_ids"].add(sid)
            entry = {}
            for h, idx in hdr.items():
                k = byh.get(h)
                if k in MANUAL_KEYS:
                    entry[k] = r[idx] if idx < len(r) else None
            prev["manual"][sid] = entry
            if "First listed here" in hdr and r[hdr["First listed here"]]:
                prev["first_seen"][sid] = r[hdr["First listed here"]]
    if "_Data" in wb.sheetnames:
        ws = wb["_Data"]
        hdr = [c.value for c in ws[1]]
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r and r[0]:
                prev["snapshot"][str(r[0])] = dict(zip(hdr, r))
    if "_Meta" in wb.sheetnames:
        ws = wb["_Meta"]
        for r in ws.iter_rows(min_row=1, values_only=True):
            if r and r[0] == "log":
                prev["log"].append(list(r[1:5]))
            elif r and r[0]:
                prev[str(r[0])] = r[1]
    return prev


# ------------------------------------------------------------------ changes
def diff(cases, prev):
    if not prev["snapshot"]:
        return None
    out = []
    now = {c["sale_id"]: c for c in cases}
    label = {"ok": "OK to Sell", "almost": "Almost Ready", "excluded": "not on sheet"}
    for sid, c in now.items():
        p = prev["snapshot"].get(sid)
        pb = p["bucket"] if p else None
        b = c["bucket"]
        if b == "excluded" and pb in (None, "excluded"):
            continue
        if pb in (None, "excluded") and b == "ok":
            out.append((1, "NEW – OK to Sell", c, "", c["status_note"]))
        elif pb == "almost" and b == "ok":
            out.append((1, "MOVED UP – now OK to Sell", c, "Almost Ready", c["status_note"]))
        elif pb in (None, "excluded") and b == "almost":
            out.append((2, "NEW – Almost Ready", c, "", c["status_note"]))
        elif pb == "ok" and b == "almost":
            out.append((2, "MOVED DOWN – now Almost Ready", c, "OK to Sell", c["status_note"]))
        elif pb in ("ok", "almost") and b == "excluded":
            out.append((0, "DROPPED", c, label[pb], c["bucket_reason"]))
        if pb == b and b != "excluded" and p:
            try:
                pob = float(p.get("opening_bid") or 0)
            except (TypeError, ValueError):
                pob = 0
            if abs(pob - float(c.get("opening_bid") or 0)) > 0.005:
                out.append((3, "BID CHANGED", c, f"${pob:,.2f}", f"now ${float(c.get('opening_bid') or 0):,.2f}"))
            new_flags = set(c["flags"]) - set((p.get("flags") or "").split("\n"))
            for fl in sorted(new_flags):
                out.append((3, "NEW FLAG", c, "", fl))
    for sid, p in prev["snapshot"].items():
        if sid not in now and p.get("bucket") in ("ok", "almost"):
            out.append((0, "DROPPED", {"sale_id": sid, "address": p.get("address"), "municipality": p.get("municipality")},
                        label[p["bucket"]], "No longer on the Sheriff's list"))
    out.sort(key=lambda x: (x[0], x[2].get("municipality") or "", x[2].get("address") or ""))
    return out


# ------------------------------------------------------------------ writers
def style_header(ws, row, ncol, manual_cols=()):
    for i in range(1, ncol + 1):
        c = ws.cell(row=row, column=i)
        c.font = WHITE_B
        c.fill = YELLOW_H if i in manual_cols else NAVY
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[row].height = 42


def write_case_tab(ws, title, subtitle, rows, status_header, today, prev):
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = subtitle
    ws["A2"].font = Font(italic=True, color="555555")
    manual_cols = []
    for i, (h, k, w, kind) in enumerate(COLUMNS, 1):
        ws.cell(row=HEADER_ROW, column=i, value=status_header if k == "status_note" else h)
        ws.column_dimensions[get_column_letter(i)].width = w
        if kind.startswith("manual"):
            manual_cols.append(i)
    style_header(ws, HEADER_ROW, len(COLUMNS), manual_cols)
    tm, av = col_of("tax_municipality"), col_of("assessed_value")
    rate_col = get_column_letter([i for i, c in enumerate(COLUMNS, 1) if c[1] == "=RATE"][0])
    for r_i, c in enumerate(rows, FIRST):
        m = prev["manual"].get(c["sale_id"], {})
        for i, (h, k, w, kind) in enumerate(COLUMNS, 1):
            cell = ws.cell(row=r_i, column=i)
            if k == "=RATE":
                cell.value = (f"=IFERROR(VLOOKUP({tm}{r_i},'Transfer Tax'!$A${TT_FIRST}:$H$400,8,FALSE),\"\")")
            elif k == "=TAX":
                cell.value = f"=IF(OR({av}{r_i}=\"\",{rate_col}{r_i}=\"\"),\"\",{av}{r_i}*'Transfer Tax'!$B$6*{rate_col}{r_i})"
            elif kind.startswith("manual"):
                cell.value = m.get(k)
                cell.fill = YELLOW
            elif k == "first_seen":
                cell.value = prev["first_seen"].get(c["sale_id"]) or today
            else:
                cell.value = c.get(k)
            if kind.endswith("money"):
                cell.number_format = MONEY_FMT
            elif kind.endswith("pct"):
                cell.number_format = PCT_FMT
            cell.alignment = Alignment(wrap_text=kind == "wrap" or k in ("flags_text",), vertical="top")
            cell.border = Border(bottom=THIN)
            if k == "status_note" and str(c.get("status_note", "")).startswith("VERIFY"):
                cell.fill = AMBER
    ws.freeze_panes = ws.cell(row=FIRST, column=4)
    last = get_column_letter(len(COLUMNS))
    ws.auto_filter.ref = f"A{HEADER_ROW}:{last}{max(FIRST, FIRST + len(rows) - 1)}"
    if not rows:
        ws.cell(row=FIRST, column=1, value="No cases in this group on the current list.")


def write_overview(ws, sale_date, source, cases, changes, prev, run_time, dropped_notes):
    ok = [c for c in cases if c["bucket"] == "ok"]
    al = [c for c in cases if c["bucket"] == "almost"]
    ws.column_dimensions["A"].width = 44
    for col, w in zip("BCDEF", (12, 34, 18, 22, 60)):
        ws.column_dimensions[col].width = w
    ws["A1"] = f"Allegheny County Sheriff Sale — {sale_date:%A, %B %-d, %Y}"
    ws["A1"].font = Font(bold=True, size=15)
    ws["A2"] = f"Source: {source}.  Updated {run_time:%a %b %-d, %Y %-I:%M %p}."
    ws["A2"].font = Font(italic=True, color="555555")
    ws["A4"], ws["B4"] = "OK to Sell", len(ok)
    ws["A5"], ws["B5"] = "Almost Ready (needs only 3129.3 / verify)", len(al)
    ws["A6"], ws["B6"] = "Residential cases still active but not ready", sum(
        1 for c in cases if c["bucket"] == "excluded" and c["residential"] and str(c.get("status", "")).lower().startswith("active"))
    ws["A7"], ws["B7"] = "All cases on the Sheriff's list", len(cases)
    for r in range(4, 8):
        ws[f"A{r}"].font = Font(bold=r < 6)
    ws["C4"] = f"Sum of opening bids: ${sum(float(c.get('opening_bid') or 0) for c in ok):,.0f}"
    ws["C5"] = f"Sum of opening bids: ${sum(float(c.get('opening_bid') or 0) for c in al):,.0f}"
    r = 9
    ws.cell(row=r, column=1, value="Changes since the last update" + (
        f" ({prev.get('last_run')})" if prev.get("last_run") else "")).font = Font(bold=True, size=12)
    r += 1
    if changes is None:
        ws.cell(row=r, column=1, value="First version for this auction — future updates will list what changed here.")
        r += 2
    elif not changes:
        ws.cell(row=r, column=1, value="No changes to OK to Sell or Almost Ready.")
        r += 2
    else:
        for i, h in enumerate(["Change", "Sale ID", "Address", "Municipality", "Was", "Detail"], 1):
            ws.cell(row=r, column=i, value=h)
        style_header(ws, r, 6)
        ws.row_dimensions[r].height = 18
        r += 1
        for pri, kind, c, was, detail in changes:
            vals = [kind, c.get("sale_id"), c.get("address"), c.get("municipality"), was, detail]
            for i, v in enumerate(vals, 1):
                cell = ws.cell(row=r, column=i, value=v)
                cell.alignment = Alignment(wrap_text=i == 6, vertical="top")
                cell.fill = GREEN if kind.startswith(("NEW –", "MOVED UP")) else AMBER if kind.startswith(("DROPPED", "MOVED DOWN")) else GREY
            r += 1
        r += 1
    if dropped_notes:
        ws.cell(row=r, column=1, value="Off the sheet, but you had entries for them (kept — they come back if the case returns)").font = Font(bold=True, size=12)
        r += 1
        for i, h in enumerate(["Sale ID", "Address", "Why it's off", "Your notes", "Max bid", ""], 1):
            ws.cell(row=r, column=i, value=h)
        style_header(ws, r, 5)
        ws.row_dimensions[r].height = 18
        r += 1
        for sid, addr, why, m in dropped_notes:
            for i, v in enumerate([sid, addr, why, m.get("my_notes"), m.get("max_bid")], 1):
                ws.cell(row=r, column=i, value=v).alignment = Alignment(wrap_text=True, vertical="top")
            r += 1
        r += 1
    ws.cell(row=r, column=1, value="Update log").font = Font(bold=True, size=12)
    r += 1
    for i, h in enumerate(["Updated", "OK to Sell", "Source file", "Almost Ready"], 1):
        ws.cell(row=r, column=i, value=h)
    style_header(ws, r, 4)
    ws.row_dimensions[r].height = 18
    r += 1
    for row in prev["log"] + [[f"{run_time:%Y-%m-%d %H:%M}", len(ok), source, len(al)]]:
        for i, v in enumerate(row, 1):
            ws.cell(row=r, column=i, value=v)
        r += 1
    r += 1
    ws.cell(row=r, column=1, value=(
        "How to use: type only in the yellow YOUR ENTRY columns on OK to Sell / Almost Ready (assessed value fills in the "
        "transfer-tax estimate). Every update keeps what you typed, matched by Sale ID. The Sheriff finalizes the list the "
        "morning of the sale — check the morning-of list before bidding. Not legal advice; see Glossary.")).alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    ws.row_dimensions[r].height = 48


def write_transfer_tax(ws, tt):
    ws["A1"] = "Realty transfer tax — every Allegheny County municipality"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = tt["source_note"]
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A2:K2")
    ws.row_dimensions[2].height = 30
    ws["A4"] = "INPUTS"
    ws["A4"].font = Font(bold=True)
    ws["A5"], ws["B5"], ws["C5"] = "State realty transfer tax rate", tt["state_rate"], "PA Dept. of Revenue — state RTT is 1%."
    ws["A6"], ws["B6"], ws["C6"] = "Allegheny CLR factor", tt["clr"], tt["clr_note"]
    ws["A7"], ws["B7"], ws["C7"] = "Example assessed value (try your own)", 40000, "Look up a parcel's assessed value on the County Real Estate portal using the 'County parcel ID' column."
    for a in ("B6", "B7"):
        ws[a].fill = YELLOW
    ws["B5"].number_format = PCT_FMT
    ws["B7"].number_format = MONEY_FMT
    ws["A9"] = "HOW THE TAX IS FIGURED ON A SHERIFF DEED"
    ws["A9"].font = Font(bold=True)
    ws["A10"] = tt["method_note"]
    ws["A10"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A10:K10")
    ws.row_dimensions[10].height = 45
    ws["A11"], ws["B11"] = "Example computed value", "=B7*B6"
    ws["B11"].number_format = MONEY_FMT
    hdr = ["Municipality", "Municipal rate", "School district", "School rate", "Local total", "State rate", "Note",
           "Total rate", "Tax per $10k assessed", "On OK to Sell", "On Almost Ready"]
    hr = TT_FIRST - 1
    for i, h in enumerate(hdr, 1):
        ws.cell(row=hr, column=i, value=h)
    style_header(ws, hr, len(hdr))
    tmc = col_of("tax_municipality")
    for r, m in enumerate(tt["municipalities"], TT_FIRST):
        vals = [m["municipality"], m["muni_rate"], m["school_district"], m["school_rate"], f"=B{r}+D{r}", "=$B$5",
                m.get("note"), f"=E{r}+F{r}", f"=10000*$B$6*H{r}",
                f"=COUNTIF('OK to Sell'!${tmc}:${tmc},A{r})", f"=COUNTIF('Almost Ready'!${tmc}:${tmc},A{r})"]
        for i, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=i, value=v)
            if i in (2, 4, 5, 6, 8):
                c.number_format = PCT_FMT
            if i == 9:
                c.number_format = MONEY_FMT
    for col, w in zip("ABCDEFGHIJK", (30, 10, 32, 10, 10, 9, 40, 10, 12, 11, 12)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = f"B{TT_FIRST}"
    ws.auto_filter.ref = f"A{hr}:K{TT_FIRST + len(tt['municipalities']) - 1}"


def write_glossary(ws, rows):
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 110
    for r, (a, b) in enumerate(rows, 1):
        ws.cell(row=r, column=1, value=a or None)
        ws.cell(row=r, column=2, value=b or None).alignment = Alignment(wrap_text=True, vertical="top")
        if a and not b:
            ws.cell(row=r, column=1).font = Font(bold=True, size=14 if r == 1 else 11)
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A2:B2")
    ws.row_dimensions[2].height = 30


# ------------------------------------------------------------------ main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("classified")
    ap.add_argument("out")
    ap.add_argument("--sale-date", required=True)
    ap.add_argument("--source", required=True)
    ap.add_argument("--previous")
    ap.add_argument("--now", help="override run time (ISO), for testing")
    a = ap.parse_args()
    sale_date = dt.date.fromisoformat(a.sale_date)
    run_time = dt.datetime.fromisoformat(a.now) if a.now else dt.datetime.now()
    today = f"{run_time:%Y-%m-%d}"
    cases = json.load(open(a.classified))
    prev = read_previous(a.previous)

    for c in cases:
        c["sale_id"] = str(c["sale_id"])
        c["boxes"] = " ".join(tick(c.get(k)) for k in ("box_service", "box_3129_2", "box_3129_3", "box_ok"))
        c["flags_text"] = "\n".join(c.get("flags") or [])
        if c["bucket"] == "almost":
            extra = [o for o in c.get("outstanding", []) if o not in c["bucket_reason"]]
            c["status_note"] = c["bucket_reason"] + ("\n• " + "\n• ".join(extra) if extra else "")
        else:
            c["status_note"] = c["bucket_reason"]

    key = lambda c: (c.get("municipality") or "", c.get("address") or "")
    ok = sorted([c for c in cases if c["bucket"] == "ok"], key=key)
    al = sorted([c for c in cases if c["bucket"] == "almost"], key=key)
    changes = diff(cases, prev)

    shown = {c["sale_id"] for c in ok + al}
    by_id = {c["sale_id"]: c for c in cases}
    dropped_notes = []
    for sid, m in prev["manual"].items():
        if sid not in shown and any(v not in (None, "") for v in m.values()):
            c = by_id.get(sid) or prev["snapshot"].get(sid) or {}
            dropped_notes.append((sid, c.get("address"), c.get("bucket_reason") or "No longer on the Sheriff's list", m))

    wb = openpyxl.Workbook()
    ov = wb.active
    ov.title = "Overview"
    write_overview(ov, sale_date, a.source, cases, changes, prev, run_time, dropped_notes)
    sd = f"{sale_date:%a %b %-d}"
    write_case_tab(wb.create_sheet("OK to Sell"), f"OK to Sell — {sd}",
                   "Active residential cases with the Sheriff's OK box checked and nothing outstanding except possibly Rule 3129.3. Status can still change before the sale.",
                   ok, "Readiness note", today, prev)
    write_case_tab(wb.create_sheet("Almost Ready"), f"Almost Ready — {sd}",
                   "Active residential cases where Service and 3129.2 are done and only Rule 3129.3 is still needed (often filed at the last minute), plus VERIFY cases where the OK box and notes disagree.",
                   al, "What's still missing", today, prev)
    write_transfer_tax(wb.create_sheet("Transfer Tax"), json.load(open(REF / "transfer_tax.json")))
    write_glossary(wb.create_sheet("Glossary"), json.load(open(REF / "glossary.json")))

    # hidden state used by the next update
    d = wb.create_sheet("_Data")
    snap_cols = ["sale_id", "case_no", "address", "municipality", "status", "bucket", "bucket_reason", "opening_bid", "flags"]
    d.append(snap_cols)
    for c in cases:
        d.append([c.get("sale_id"), c.get("case_no"), c.get("address"), c.get("municipality"), c.get("status"),
                  c["bucket"], c["bucket_reason"], c.get("opening_bid"), c["flags_text"]])
    mstore = wb.create_sheet("_Manual")
    mstore.append(["sale_id"] + MANUAL_KEYS)
    merged = dict(prev["manual"])
    for sid, m in merged.items():
        if any(v not in (None, "") for v in m.values()):
            mstore.append([sid] + [m.get(k) for k in MANUAL_KEYS])
    meta = wb.create_sheet("_Meta")
    meta.append(["sale_date", a.sale_date])
    meta.append(["source", a.source])
    meta.append(["last_run", f"{run_time:%Y-%m-%d %H:%M}"])
    for row in prev["log"] + [[f"{run_time:%Y-%m-%d %H:%M}", len(ok), a.source, len(al)]]:
        meta.append(["log"] + row)
    for s in (d, mstore, meta):
        s.sheet_state = "hidden"
    wb.save(a.out)
    print(f"OK to Sell {len(ok)} · Almost Ready {len(al)} · changes {len(changes) if changes is not None else 'n/a (first run)'}"
          f" · kept entries for {sum(1 for m in merged.values() if any(v not in (None, '') for v in m.values()))} cases")


if __name__ == "__main__":
    main()
